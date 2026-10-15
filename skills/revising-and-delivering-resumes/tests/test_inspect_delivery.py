from __future__ import annotations

import json
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path

from docx import Document
from docx.shared import Pt
from reportlab.pdfgen import canvas


SCRIPT = Path(__file__).resolve().parents[1] / "scripts" / "inspect_delivery.py"


def make_docx(path: Path, *, anomalous: bool = False) -> None:
    document = Document()
    first = document.add_paragraph()
    run = first.add_run("Normal body phrase")
    run.font.name = "Helvetica"
    run.font.size = Pt(9)
    second = document.add_paragraph()
    for text, size in (("Anomalous ", 9), ("body phrase", 12 if anomalous else 9)):
        run = second.add_run(text)
        run.font.name = "Helvetica"
        run.font.size = Pt(size)
    document.save(path)


def make_pdf(path: Path, *, anomalous: bool) -> None:
    output = canvas.Canvas(str(path), pagesize=(595.28, 841.89))
    output.setFont("Helvetica", 9)
    output.drawString(72, 760, "Normal body phrase")
    output.setFont("Times-Roman" if anomalous else "Helvetica", 12 if anomalous else 9)
    output.drawString(72, 742, "Anomalous body phrase")
    output.save()


def run_inspector(
    docx_path: Path, pdf_path: Path, *extra_arguments: str
) -> subprocess.CompletedProcess[str]:
    return subprocess.run(
        [
            sys.executable,
            str(SCRIPT),
            "--review-docx",
            str(docx_path),
            "--application-pdf",
            str(pdf_path),
            "--paper",
            "a4",
            "--forbid-pdf-typography-drift",
            *extra_arguments,
        ],
        check=False,
        capture_output=True,
        text=True,
    )


class InspectDeliveryTypographyTest(unittest.TestCase):
    def test_rejects_pdf_font_and_size_drift_from_word(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            docx_path = root / "resume.docx"
            pdf_path = root / "resume.pdf"
            make_docx(docx_path)
            make_pdf(pdf_path, anomalous=True)

            result = run_inspector(docx_path, pdf_path)

            self.assertEqual(result.returncode, 1, result.stdout + result.stderr)
            report = json.loads(result.stdout)
            check = next(item for item in report["checks"] if item["name"] == "pdf_typography_matches_word")
            self.assertFalse(check["passed"])
            self.assertIn("Anomalous body phrase", check["detail"])

    def test_accepts_pdf_typography_matching_word(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            docx_path = root / "resume.docx"
            pdf_path = root / "resume.pdf"
            make_docx(docx_path)
            make_pdf(pdf_path, anomalous=False)

            result = run_inspector(docx_path, pdf_path)

            self.assertEqual(result.returncode, 0, result.stdout + result.stderr)
            report = json.loads(result.stdout)
            check = next(item for item in report["checks"] if item["name"] == "pdf_typography_matches_word")
            self.assertTrue(check["passed"])

    def test_rejects_mixed_body_font_sizes_inside_word_paragraph(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            docx_path = root / "resume.docx"
            pdf_path = root / "resume.pdf"
            make_docx(docx_path, anomalous=True)
            make_pdf(pdf_path, anomalous=False)

            result = run_inspector(docx_path, pdf_path, "--forbid-word-mixed-font-sizes")

            self.assertEqual(result.returncode, 1, result.stdout + result.stderr)
            report = json.loads(result.stdout)
            check = next(item for item in report["checks"] if item["name"] == "word_paragraph_font_sizes_consistent")
            self.assertFalse(check["passed"])
            self.assertIn("Anomalous body phrase", check["detail"])


if __name__ == "__main__":
    unittest.main()
